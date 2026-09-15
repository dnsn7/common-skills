#!/usr/bin/env python
import argparse
import json
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def font(run, name="Microsoft YaHei", size=11, bold=False, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def main():
    ap = argparse.ArgumentParser(description="Build standalone Chinese slide narration from notes JSON.")
    ap.add_argument("notes_json")
    ap.add_argument("--out", required=True)
    ap.add_argument("--title", default="论文组会汇报演讲稿")
    ap.add_argument("--date", default="")
    ap.add_argument("--presenter", default="汇报人")
    ap.add_argument("--institution", default="研究生组会")
    args = ap.parse_args()
    records = json.loads(Path(args.notes_json).read_text(encoding="utf-8"))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for key in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, key, Inches(1))
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13)):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("2E74B5")
    header = sec.header.paragraphs[0]
    header.text = args.institution
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        font(run, size=9, color="5B7083")
    p = doc.add_paragraph()
    font(p.add_run(args.title), size=24, bold=True, color="102A43")
    identity = "　".join(part for part in (f"汇报人：{args.presenter}", args.institution, args.date) if part)
    p = doc.add_paragraph(identity)
    for run in p.runs:
        font(run, bold=True)
    for index, record in enumerate(records, 1):
        doc.add_heading(f"第 {index} 页　{record.get('title', '')}".rstrip(), level=1)
        doc.add_paragraph(record.get("body", ""))
        sources = record.get("sources", [])
        if sources:
            p = doc.add_paragraph()
            font(p.add_run("来源："), bold=True, color="2563A6")
            for source in sources:
                p = doc.add_paragraph(source, style="List Bullet")
                p.paragraph_format.left_indent = Inches(0.38)
                p.paragraph_format.first_line_indent = Inches(-0.19)
                p.paragraph_format.space_after = Pt(4)
    out = Path(args.out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
